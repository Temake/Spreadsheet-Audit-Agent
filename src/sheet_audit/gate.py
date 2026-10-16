"""Gate 1: decide whether a sheet is a table we can audit, and shape it into a DataFrame.

Refusing early ("no header row found") is safer than auditing something that is not a
table and reporting anomalies that are really just layout.
"""

import math
from dataclasses import dataclass

import pandas as pd
from openpyxl.utils import get_column_letter

from sheet_audit.config import Settings
from sheet_audit.sanitize import clean
from sheet_audit.schemas import GateResult
from sheet_audit.values import looks_like_label, normalize_cell

HEADER_SCAN_ROWS = 10


@dataclass
class SheetTable:
    name: str
    # Cells are str or None. The index is the 0-based raw row position, so the
    # spreadsheet row number the user sees is index + 1.
    df: pd.DataFrame
    gate: GateResult


def row_number(index: int) -> int:
    return int(index) + 1


def _find_header(rows: list[list], width: int) -> int | None:
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        filled = [v for v in row if v is not None]
        if len(filled) < max(2, math.ceil(0.5 * width)):
            continue
        if sum(looks_like_label(v) for v in filled) / len(filled) >= 0.8:
            return i
    return None


def _fail(name: str, reason: str, **kwargs) -> SheetTable:
    gate = GateResult(sheet=name, auditable=False, reasons=[reason], **kwargs)
    return SheetTable(name=name, df=pd.DataFrame(), gate=gate)


def prepare_sheet(name: str, raw_rows: list[list], settings: Settings) -> SheetTable:
    name = clean(name, 64) or "Sheet"
    rows = [[normalize_cell(v) for v in r] for r in raw_rows]
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    if not rows:
        return _fail(name, "sheet is empty")

    width = max(len(r) for r in rows)
    rows = [r + [None] * (width - len(r)) for r in rows]
    densest = max(sum(v is not None for v in r) for r in rows[:50])
    header_idx = _find_header(rows, densest)
    if header_idx is None:
        return _fail(name, f"no header row found in the first {HEADER_SCAN_ROWS} rows")

    warnings: list[str] = []
    if header_idx > 0:
        warnings.append(f"skipped {header_idx} row(s) above the header (titles or notes)")

    data = rows[header_idx + 1 :]
    df = pd.DataFrame(data, index=range(header_idx + 1, header_idx + 1 + len(data)), dtype=object)
    header = rows[header_idx]

    names, seen = [], set()
    keep = []
    for col in range(width):
        label = header[col]
        if label is None and df[col].isna().all():
            continue  # spacer column
        if label is None:
            label = f"column_{get_column_letter(col + 1)}"
            warnings.append(f"column {get_column_letter(col + 1)} has data but no header; named {label!r}")
        label = clean(label, 128)
        base, n = label, 2
        while label in seen:
            label = f"{base}_{n}"
            n += 1
        if label != base:
            warnings.append(f"duplicate header {base!r} renamed to {label!r}")
        seen.add(label)
        names.append(label)
        keep.append(col)
    df = df[keep]
    df.columns = names

    blank = df.isna().all(axis=1)
    if blank.any():
        warnings.append(f"ignored {int(blank.sum())} blank row(s) inside the data; the sheet may hold more than one table")
        df = df[~blank]

    gate = GateResult(
        sheet=name,
        auditable=True,
        warnings=warnings,
        header_row=header_idx + 1,
        n_rows=len(df),
        n_columns=len(df.columns),
    )
    if len(df) < settings.min_rows:
        gate.auditable = False
        gate.reasons.append(f"only {len(df)} data row(s); at least {settings.min_rows} are needed")
    if len(df.columns) < settings.min_columns:
        gate.auditable = False
        gate.reasons.append(f"only {len(df.columns)} column(s); at least {settings.min_columns} are needed")
    if len(df) > settings.max_rows:
        gate.auditable = False
        gate.reasons.append(f"more than {settings.max_rows} data rows")
    return SheetTable(name=name, df=df, gate=gate)
